from __future__ import annotations

import io
import json
import sqlite3
import uuid
from datetime import datetime
from pathlib import Path

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from audit_engine import run_checks


ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
DATA.mkdir(exist_ok=True)
DB_PATH = DATA / "review_feedback.db"

app = FastAPI(title="视频文字检查工具")
app.mount("/static", StaticFiles(directory=ROOT / "static"), name="static")
tasks: dict[str, dict] = {}


def init_db() -> None:
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS feedback (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT NOT NULL,
                task_id TEXT NOT NULL,
                issue_id TEXT NOT NULL,
                video TEXT NOT NULL,
                issue_kind TEXT NOT NULL,
                sentence TEXT NOT NULL,
                review_result TEXT NOT NULL,
                note TEXT NOT NULL
            )
        """)


init_db()


@app.get("/", response_class=HTMLResponse)
async def home():
    return (ROOT / "templates" / "index.html").read_text(encoding="utf-8")


@app.post("/api/check")
async def check_files(
    subtitles: list[UploadFile] = File(...),
    lexicon: UploadFile = File(...),
    document: UploadFile | None = File(None),
    video_names: str = Form("[]"),
):
    if not subtitles:
        raise HTTPException(400, "请至少选择一个字幕文件")
    srt_files = [(item.filename or "未命名字幕.srt", await item.read()) for item in subtitles]
    lexicon_data = await lexicon.read()
    document_data = await document.read() if document else None
    issues, summary = run_checks(srt_files, document_data, lexicon_data)
    task_id = uuid.uuid4().hex[:12]
    tasks[task_id] = {
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "video_names": json.loads(video_names or "[]"),
        "issues": issues,
        "summary": summary,
        "lexicon_name": lexicon.filename or "审核词表.xlsx",
    }
    return {"task_id": task_id, "issues": issues, "summary": summary, "lexicon_name": tasks[task_id]["lexicon_name"]}


@app.post("/api/tasks/{task_id}/review")
async def save_review(task_id: str, payload: dict):
    task = tasks.get(task_id)
    if not task:
        raise HTTPException(404, "本次检查已清除，请重新上传文件")
    updates = {str(item.get("id")): item for item in payload.get("issues", [])}
    with sqlite3.connect(DB_PATH) as conn:
        for row in task["issues"]:
            if row["id"] not in updates:
                continue
            update = updates[row["id"]]
            row["review_result"] = str(update.get("review_result", ""))
            row["note"] = str(update.get("note", ""))
            if row["review_result"]:
                conn.execute(
                    "INSERT INTO feedback(created_at,task_id,issue_id,video,issue_kind,sentence,review_result,note) VALUES(?,?,?,?,?,?,?,?)",
                    (datetime.now().isoformat(timespec="seconds"), task_id, row["id"], row["video"], row["kind"], row["sentence"], row["review_result"], row["note"]),
                )
    return {"ok": True}


@app.get("/api/tasks/{task_id}/report.xlsx")
async def download_report(task_id: str):
    task = tasks.get(task_id)
    if not task:
        raise HTTPException(404, "本次检查已清除，请重新上传文件")
    book = Workbook()
    sheet = book.active
    sheet.title = "问题清单"
    headers = ["编号", "视频", "涉及时间", "发现了什么", "需要查看的完整句子", "问题说清楚", "建议怎么处理", "系统初步结果", "你的复核结果", "你的备注"]
    sheet.append(headers)
    for idx, row in enumerate(task["issues"], 1):
        sheet.append([idx, row["video"], row["time"], row["kind"], row["sentence"], row["detail"], row["suggestion"], row["system_result"], row["review_result"], row["note"]])
    header_fill = PatternFill("solid", fgColor="24577A")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = [8, 32, 22, 20, 38, 55, 42, 18, 18, 35]
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + col)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    output = io.BytesIO()
    book.save(output)
    output.seek(0)
    headers_out = {"Content-Disposition": f'attachment; filename="video-review-{task_id}.xlsx"'}
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers_out)


@app.get("/api/improvement-suggestions.xlsx")
async def download_suggestions():
    book = Workbook()
    sheet = book.active
    sheet.title = "待整理建议"
    sheet.append(["记录时间", "视频", "系统提示类型", "句子", "审核人员的选择", "审核人员备注", "建议怎么改规则"])
    with sqlite3.connect(DB_PATH) as conn:
        rows = conn.execute("SELECT created_at,video,issue_kind,sentence,review_result,note FROM feedback ORDER BY id DESC").fetchall()
    for created, video, kind, sentence, result, note in rows:
        suggestion = ""
        if result == "系统提示错了": suggestion = "考虑增加允许写法，避免下次重复提示"
        elif result == "确实有问题": suggestion = "检查是否需要补充词表或检查规则"
        elif result == "已经处理好": suggestion = "把这种正确处理方式加入固定测试样本"
        elif result == "我不确定": suggestion = "请熟悉审核标准的人再次确认"
        sheet.append([created, video, kind, sentence, result, note, suggestion])
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="24577A")
        cell.font = Font(color="FFFFFF", bold=True)
    for col, width in enumerate([22, 32, 22, 45, 20, 40, 42], 1):
        sheet.column_dimensions[chr(64 + col)].width = width
    for row in sheet.iter_rows():
        for cell in row: cell.alignment = Alignment(vertical="top", wrap_text=True)
    output = io.BytesIO(); book.save(output); output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="improvement-suggestions.xlsx"'})


@app.delete("/api/tasks/{task_id}")
async def clear_task(task_id: str):
    tasks.pop(task_id, None)
    return {"ok": True}


@app.get("/favicon.ico", include_in_schema=False)
async def favicon():
    return FileResponse(ROOT / "static" / "favicon.svg", media_type="image/svg+xml")
