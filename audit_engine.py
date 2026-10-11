from __future__ import annotations

import io
import re
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


TIME_RE = re.compile(r"(\d{2}):(\d{2}):(\d{2})[,.](\d{3})")


@dataclass
class Cue:
    number: int
    start_ms: int
    end_ms: int
    text: str


def time_to_ms(value: str) -> int:
    match = TIME_RE.search(value)
    if not match:
        return 0
    h, m, s, ms = map(int, match.groups())
    return ((h * 60 + m) * 60 + s) * 1000 + ms


def format_time(value: int) -> str:
    value = max(0, value)
    h, rest = divmod(value, 3_600_000)
    m, rest = divmod(rest, 60_000)
    s, ms = divmod(rest, 1000)
    return f"{h:02d}:{m:02d}:{s:02d}.{ms:03d}"


def parse_srt(data: bytes) -> list[Cue]:
    text = data.decode("utf-8-sig", errors="replace").replace("\r\n", "\n")
    cues: list[Cue] = []
    for block in re.split(r"\n\s*\n", text.strip()):
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        if len(lines) < 2:
            continue
        number = int(lines[0]) if lines[0].isdigit() else len(cues) + 1
        time_index = 1 if "-->" in lines[1] else 0
        if "-->" not in lines[time_index]:
            continue
        start, end = [part.strip() for part in lines[time_index].split("-->", 1)]
        cue_text = " ".join(lines[time_index + 1 :]).strip()
        cues.append(Cue(number, time_to_ms(start), time_to_ms(end), cue_text))
    return cues


def normalize(text: str) -> str:
    text = re.sub(r"【提炼】|宇廷老师：|学员：|Q[:：]", "", text)
    return "".join(re.findall(r"[\u4e00-\u9fffA-Za-z0-9]+", text)).lower()


def read_docx(data: bytes) -> list[str]:
    doc = Document(io.BytesIO(data))
    return [p.text.strip() for p in doc.paragraphs if p.text.strip()]


def load_terms(data: bytes) -> tuple[list[dict], str]:
    book = load_workbook(io.BytesIO(data), read_only=False, data_only=True)
    sheet = book["敏感词库"] if "敏感词库" in book.sheetnames else book.worksheets[0]
    header_row = None
    headers: list[str] = []
    for row_no, row in enumerate(sheet.iter_rows(values_only=True), 1):
        values = [str(v).strip() if v is not None else "" for v in row]
        if "敏感词" in values:
            header_row, headers = row_no, values
            break
    if not header_row:
        return [], "未找到“敏感词”这一列"
    index = {name: i for i, name in enumerate(headers) if name}
    terms = []
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        values = list(row)
        word = str(values[index["敏感词"]]).strip() if index["敏感词"] < len(values) and values[index["敏感词"]] else ""
        enabled = str(values[index.get("是否启用", -1)]).strip() if index.get("是否启用", -1) >= 0 and index.get("是否启用", -1) < len(values) and values[index.get("是否启用", -1)] else "是"
        if not word or enabled not in {"是", "启用", "1", "True", "true"}:
            continue
        terms.append({
            "word": word,
            "category": str(values[index.get("分类", -1)] or "") if index.get("分类", -1) >= 0 else "",
            "suggestion": str(values[index.get("建议处理", -1)] or "") if index.get("建议处理", -1) >= 0 else "",
        })
    terms.sort(key=lambda item: len(item["word"]), reverse=True)
    return terms, Path("审核词表.xlsx").stem


def issue(issue_id: str, video: str, time_range: str, sentence: str, kind: str, detail: str, suggestion: str, system_result: str) -> dict:
    return {
        "id": issue_id,
        "video": video,
        "time": time_range,
        "sentence": sentence,
        "kind": kind,
        "detail": detail,
        "suggestion": suggestion,
        "system_result": system_result,
        "review_result": "",
        "note": "",
    }


def find_doc_section(paragraphs: list[str], title: str) -> str:
    title_key = normalize(title)
    start = 0
    for i, paragraph in enumerate(paragraphs):
        p = normalize(paragraph)
        if title_key and (title_key in p or SequenceMatcher(None, title_key, p).ratio() > 0.72):
            start = i
            break
    selected = []
    for paragraph in paragraphs[start:]:
        if selected and re.match(r"^\d+[、.．]", paragraph):
            break
        if paragraph.startswith("学员：") and "【提炼】" not in paragraph:
            continue
        selected.append(paragraph)
    return normalize("".join(selected))


def run_checks(srt_files: list[tuple[str, bytes]], docx_data: bytes | None, lexicon_data: bytes) -> tuple[list[dict], dict]:
    terms, _ = load_terms(lexicon_data)
    paragraphs = read_docx(docx_data) if docx_data else []
    issues: list[dict] = []
    summary = {"files": len(srt_files), "terms": len(terms), "order": 0, "overlap": 0, "word_hits": 0, "text_differences": 0}

    for file_no, (filename, raw) in enumerate(srt_files, 1):
        title = Path(filename).stem
        cues = parse_srt(raw)
        for pos, cue in enumerate(cues):
            if pos and cue.start_ms < cues[pos - 1].start_ms:
                prev = cues[pos - 1]
                detail = (
                    f"字幕文件第{cue.number}条，时间是{format_time(cue.start_ms)}–{format_time(cue.end_ms)}，"
                    f"却排在第{prev.number}条（{format_time(prev.start_ms)}–{format_time(prev.end_ms)}）后面。"
                    "按照时间，这句话应该移到前面相应位置。"
                )
                issues.append(issue(f"{file_no}-order-{cue.number}", title, f"{format_time(cue.start_ms)}–{format_time(cue.end_ms)}", cue.text, "字幕位置放错", detail, "按开始时间重新排列字幕，并重新编号。", "需要修改"))
                summary["order"] += 1

        ordered = sorted(cues, key=lambda c: (c.start_ms, c.end_ms))
        for previous, current in zip(ordered, ordered[1:]):
            if current.start_ms < previous.end_ms:
                overlap = previous.end_ms - current.start_ms
                detail = (
                    f"前一句“{previous.text}”到{format_time(previous.end_ms)}结束；"
                    f"后一句“{current.text}”从{format_time(current.start_ms)}开始，两句重叠{overlap / 1000:.3f}秒。"
                )
                issues.append(issue(f"{file_no}-overlap-{current.number}", title, f"{format_time(current.start_ms)}–{format_time(current.end_ms)}", current.text, "两句字幕时间重叠", detail, "调整相接处的开始或结束时间，确保两句不重叠。", "需要修改"))
                summary["overlap"] += 1

        for cue in cues:
            occupied: list[tuple[int, int]] = []
            for term in terms:
                word = term["word"]
                for match in re.finditer(re.escape(word), cue.text, re.IGNORECASE):
                    span = match.span()
                    if any(span[0] >= a and span[1] <= b for a, b in occupied):
                        continue
                    occupied.append(span)
                    short_note = "这是单字或很短的词，请结合整句话确认。" if len(word) <= 1 else ""
                    detail = f"审核词表中有“{word}”。{short_note}"
                    suggestion = "确认需要处理后，把这个词全部改成空格，或者改成连续拼音或大写首字母。"
                    issues.append(issue(f"{file_no}-word-{cue.number}-{span[0]}", title, f"{format_time(cue.start_ms)}–{format_time(cue.end_ms)}", cue.text, f"发现词语“{word}”", detail, suggestion, "需要确认" if len(word) <= 1 else "需要修改"))
                    summary["word_hits"] += 1

        if paragraphs:
            doc_text = find_doc_section(paragraphs, title)
            for cue in cues:
                cue_text = normalize(cue.text)
                if len(cue_text) < 2 or cue_text in doc_text:
                    continue
                best = max((SequenceMatcher(None, cue_text, doc_text[i:i+len(cue_text)+8]).ratio() for i in range(0, max(1, len(doc_text)-len(cue_text)), max(1, len(cue_text)//3))), default=0)
                if best < 0.72:
                    issues.append(issue(f"{file_no}-text-{cue.number}", title, f"{format_time(cue.start_ms)}–{format_time(cue.end_ms)}", cue.text, "字幕和文稿可能不一样", "这句字幕没有在对应文稿中找到相同内容。请听原视频并查看Word文稿。", "确认字幕是否少字、错字或多字；如有问题，在备注中写出正确句子。", "需要确认"))
                    summary["text_differences"] += 1

    return issues, summary
